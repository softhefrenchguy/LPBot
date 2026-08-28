import os
from dotenv import load_dotenv

env_path = os.path.join(os.path.dirname(__file__), ".env")
load_dotenv(dotenv_path=env_path, override=True)
import os
from dotenv import load_dotenv

env_path = os.path.join(os.path.dirname(__file__), ".env")
load_dotenv(dotenv_path=env_path, override=True)
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

HISTORY_FILE = "portfolio_history.jsonl"
EXCEL_FILE = "portfolio_history.xlsx"

# Load JSONL â†’ list of dicts
with open(HISTORY_FILE) as f:
    lines = [json.loads(l) for l in f if l.strip()]

# Build dataframe
df = pd.DataFrame(lines)

# --- FIX: strip timezone ---
df["timestamp"] = pd.to_datetime(df["timestamp"], utc=True).dt.tz_localize(None)
df = df.sort_values("timestamp")

# Save to Excel
df.to_excel(EXCEL_FILE, index=False)

# Add chart
wb = load_workbook(EXCEL_FILE)
ws = wb.active

chart = LineChart()
chart.title = "Portfolio Performance Over Time"
chart.y_axis.title = "Value (USD)"
chart.x_axis.title = "Time (UTC)"

values = Reference(ws, min_col=3, max_col=6, min_row=1, max_row=len(df)+1)
cats = Reference(ws, min_col=1, min_row=2, max_row=len(df)+1)
chart.add_data(values, titles_from_data=True)
chart.set_categories(cats)

ws.add_chart(chart, "I2")
wb.save(EXCEL_FILE)

print(f"âœ… Exported {len(df)} entries â†’ {EXCEL_FILE}")
